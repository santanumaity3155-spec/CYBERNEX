import re
import os
import uuid
from typing import Dict, Any, List, Optional
from app.core.config import get_settings
from app.core.logging import logger

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import pptx
    from pptx.util import Inches as PPTXInches, Pt as PPTXPt
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

settings = get_settings()

# Match characters disallowed in XML 1.0 documents:
# Valid chars: #x9 | #xA | #xD | [#x20-#xD7FF] | [#xE000-#xFFFD] | [#x10000-#x10FFFF]
_ILLEGAL_XML_CHARS_RE = re.compile(
    r"[^\u0009\u000A\u000D\u0020-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]"
)


def sanitize_xml_text(text: Any) -> str:
    """
    Sanitizes string for XML 1.0 / OpenXML compliance.
    Replaces form feeds (\\x0c) and vertical tabs (\\x0b) with newlines,
    and removes NULL bytes (\\x00) and any other control characters
    disallowed in XML 1.0 documents.
    """
    if text is None:
        return ""
    if not isinstance(text, str):
        text = str(text)
    # Convert form feed and vertical tab to newline so text layout is preserved
    text = text.replace("\x0c", "\n").replace("\x0b", "\n")
    # Strip any characters illegal in XML 1.0
    return _ILLEGAL_XML_CHARS_RE.sub("", text)


class DocumentGenerator:
    def generate_docx(
        self,
        title: str,
        sections: List[Dict[str, str]],
        output_name: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Generates a sovereign DOCX deliverable in storage/outputs/.
        Enforces a genuine OpenXML ZIP package created via python-docx.
        Never saves plain text under a .docx extension.
        """
        if not DOCX_AVAILABLE:
            raise RuntimeError(
                "python-docx is not installed. Cannot generate valid DOCX package."
            )

        settings.init_storage_dirs()
        doc_id = f"docgen-{uuid.uuid4().hex[:8]}"
        filename = output_name or f"Deliverable_{doc_id}.docx"
        if not filename.endswith(".docx"):
            filename += ".docx"

        file_path = os.path.abspath(os.path.join(settings.OUTPUT_DIR, filename))

        clean_title = sanitize_xml_text(title)

        try:
            doc = docx.Document()

            # Header Title
            title_p = doc.add_heading(level=0)
            run = title_p.add_run(clean_title)
            run.font.color.rgb = RGBColor(12, 74, 110)  # Cybernex Sky Blue #0C4A6E
            run.font.size = Pt(22)
            run.bold = True

            doc.add_paragraph("CYBERNEX Sovereign AI Workbench Deliverable")
            doc.add_paragraph("=" * 60)

            for sec in sections:
                sec_title = sanitize_xml_text(sec.get("title", "Section"))
                sec_content = sanitize_xml_text(sec.get("content", ""))

                h = doc.add_heading(sec_title, level=1)
                if h.runs:
                    h.runs[0].font.color.rgb = RGBColor(3, 105, 161)

                if sec_content:
                    blocks = sec_content.split("\n\n")
                    for b in blocks:
                        clean_b = b.strip()
                        if not clean_b:
                            continue
                        lines = clean_b.split("\n")
                        for line in lines:
                            clean_line = line.strip()
                            if clean_line:
                                p = doc.add_paragraph(clean_line)
                                p.style.font.size = Pt(11)
                else:
                    p = doc.add_paragraph("")
                    p.style.font.size = Pt(11)

            doc.save(file_path)
            logger.info(f"Generated valid DOCX deliverable: {file_path}")

        except Exception as e:
            logger.error(f"Failed to generate valid DOCX file: {e}")
            if os.path.exists(file_path):
                try:
                    os.remove(file_path)
                except OSError:
                    pass
            raise RuntimeError(f"Failed to generate valid DOCX package: {e}") from e

        size_bytes = os.path.getsize(file_path) if os.path.exists(file_path) else 1024
        size_str = f"{round(size_bytes / 1024, 1)} KB"

        return {
            "id": doc_id,
            "name": filename,
            "type": "DOCX",
            "size": size_str,
            "status": "Verified",
            "summary": f"Generated formal document '{clean_title}'.",
            "file_path": file_path,
            "download_url": f"/api/v1/documents/{doc_id}/download"
        }

    def generate_xlsx(
        self,
        title: str,
        rows: List[List[Any]],
        output_name: Optional[str] = None
    ) -> Dict[str, Any]:
        settings.init_storage_dirs()
        doc_id = f"docgen-{uuid.uuid4().hex[:8]}"
        filename = output_name or f"Analysis_{doc_id}.xlsx"
        if not filename.endswith(".xlsx"):
            filename += ".xlsx"

        file_path = os.path.abspath(os.path.join(settings.OUTPUT_DIR, filename))

        if not OPENPYXL_AVAILABLE:
            raise RuntimeError("openpyxl is not installed. Cannot generate valid XLSX deliverable.")

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sanitize_xml_text(title)[:30]
            for r_idx, row in enumerate(rows, 1):
                for c_idx, val in enumerate(row, 1):
                    clean_val = sanitize_xml_text(val) if isinstance(val, str) else val
                    ws.cell(row=r_idx, column=c_idx, value=clean_val)
            wb.save(file_path)
        except Exception as e:
            logger.error(f"Failed to generate valid XLSX file: {e}")
            if os.path.exists(file_path):
                try:
                    os.remove(file_path)
                except OSError:
                    pass
            raise RuntimeError(f"Failed to generate valid XLSX package: {e}") from e

        size_bytes = os.path.getsize(file_path) if os.path.exists(file_path) else 1024
        return {
            "id": doc_id,
            "name": filename,
            "type": "XLSX",
            "size": f"{round(size_bytes / 1024, 1)} KB",
            "status": "Verified",
            "summary": f"Generated spreadsheet deliverable '{title}'.",
            "file_path": file_path,
            "download_url": f"/api/v1/documents/{doc_id}/download"
        }

    def _fallback_text(self, file_path: str, title: str, sections: List[Dict[str, str]]):
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(f"CYBERNEX SOVEREIGN EXECUTIVE REPORT: {title}\n\n")
            for sec in sections:
                f.write(f"=== {sec.get('title', '')} ===\n")
                f.write(f"{sec.get('content', '')}\n\n")


doc_generator = DocumentGenerator()

